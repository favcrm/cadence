#!/usr/bin/env python3
"""Bounded local extraction for text PDF, DOCX and simple XLSX sources."""

from __future__ import annotations

import argparse
from contextlib import contextmanager
import datetime
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import stat
import sys
import tempfile
import zipfile

MAX_INPUT = 20 * 1024 * 1024
MAX_PDF_PAGES = 100
MAX_EXPANDED = 100 * 1024 * 1024
MAX_XLSX_CELLS = 20_000
CHUNK = 1024 * 1024


class Refusal(Exception):
    pass


def refuse(message: str) -> None:
    raise Refusal(f"Input refused: {message}")


def stat_signature(info: os.stat_result) -> tuple[int, int, int, int, int, int]:
    return (info.st_dev, info.st_ino, info.st_mode, info.st_size, info.st_mtime_ns, info.st_ctime_ns)


def descriptor_digest(descriptor: int) -> str:
    value = hashlib.sha256()
    offset = 0
    while True:
        block = os.pread(descriptor, CHUNK, offset)
        if not block:
            return value.hexdigest()
        value.update(block)
        offset += len(block)


@contextmanager
def source_snapshot(path: Path, allowed_suffixes: set[str] | None = None):
    unresolved = path.expanduser().absolute()
    if unresolved.is_symlink() or any(parent.is_symlink() for parent in unresolved.parents):
        refuse("source path must not contain symbolic links")
    source_name = unresolved.name
    suffix = unresolved.suffix.lower()
    if suffix in {".docm", ".xlsm", ".pptm"}:
        refuse("unsupported macro-enabled Office extension")
    supported = allowed_suffixes if allowed_suffixes is not None else {".pdf", ".docx", ".xlsx"}
    if suffix not in supported:
        refuse("unsupported source extension")
    if not hasattr(os, "O_NOFOLLOW") or not hasattr(os, "pread"):
        refuse("safe no-follow source snapshot is unavailable on this platform")
    try:
        source_path = unresolved.resolve(strict=True)
        source_fd = os.open(source_path, os.O_RDONLY | os.O_NOFOLLOW)
    except (OSError, RuntimeError) as exc:
        refuse(f"cannot open source without following links ({exc})")
    temporary = None
    try:
        temporary = tempfile.TemporaryDirectory(prefix="business-intake-", dir="/tmp")
        before = os.fstat(source_fd)
        if not stat.S_ISREG(before.st_mode):
            refuse("source must be a regular file")
        if before.st_size > MAX_INPUT:
            refuse(f"input exceeds {MAX_INPUT // (1024 * 1024)} MiB limit")
        try:
            path_stat = os.stat(unresolved, follow_symlinks=False)
        except OSError as exc:
            refuse(f"source path changed while opening ({exc})")
        if stat_signature(path_stat) != stat_signature(before):
            refuse("source path was replaced while opening")

        snapshot_path = Path(temporary.name) / f"input{suffix}"
        snapshot_fd = os.open(snapshot_path, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o600)
        source_hash = hashlib.sha256()
        byte_count = 0
        try:
            with os.fdopen(snapshot_fd, "wb") as snapshot_file:
                while True:
                    block = os.read(source_fd, CHUNK)
                    if not block:
                        break
                    byte_count += len(block)
                    if byte_count > MAX_INPUT:
                        refuse(f"input exceeds {MAX_INPUT // (1024 * 1024)} MiB limit")
                    source_hash.update(block)
                    snapshot_file.write(block)
                snapshot_file.flush()
                os.fsync(snapshot_file.fileno())
        except BaseException:
            snapshot_path.unlink(missing_ok=True)
            raise

        after = os.fstat(source_fd)
        snapshot_hash = source_hash.hexdigest()
        try:
            path_after = os.stat(unresolved, follow_symlinks=False)
        except OSError as exc:
            refuse(f"source path changed during snapshot ({exc})")
        if (stat_signature(before) != stat_signature(after)
                or stat_signature(after) != stat_signature(path_after)
                or byte_count != after.st_size
                or descriptor_digest(source_fd) != snapshot_hash):
            refuse("source changed or was replaced while snapshotting")
        snapshot_path.chmod(0o400)
        state = {
            "path": unresolved,
            "resolved": source_path,
            "filename": source_name,
            "snapshot": snapshot_path,
            "sha256": snapshot_hash,
            "fd": source_fd,
            "signature": stat_signature(after),
        }
        yield state
    finally:
        os.close(source_fd)
        if temporary is not None:
            temporary.cleanup()


def verify_source_snapshot(snapshot: dict) -> None:
    try:
        descriptor_stat = os.fstat(snapshot["fd"])
        path_stat = os.stat(snapshot["path"], follow_symlinks=False)
    except OSError as exc:
        refuse(f"source changed or was replaced during extraction ({exc})")
    if (stat_signature(descriptor_stat) != snapshot["signature"]
            or stat_signature(path_stat) != snapshot["signature"]
            or descriptor_digest(snapshot["fd"]) != snapshot["sha256"]):
        refuse("source changed or was replaced during extraction")


def zip_members(source: Path) -> zipfile.ZipFile:
    try:
        package = zipfile.ZipFile(source)
        infos = package.infolist()
        if sum(item.file_size for item in infos) > MAX_EXPANDED:
            package.close()
            refuse("Office package expands beyond 100 MiB")
        if any(item.flag_bits & 1 for item in infos):
            package.close()
            refuse("encrypted Office package is unsupported")
        names = [item.filename.replace("\\", "/") for item in infos]
        if len(names) != len(set(names)) or any(name.startswith("/") or "../" in name for name in names):
            package.close()
            refuse("Office package has unsafe or duplicate paths")
        return package
    except (OSError, zipfile.BadZipFile, RuntimeError) as exc:
        refuse(f"corrupt Office package ({exc})")


def preflight_office(source: Path, kind: str) -> zipfile.ZipFile:
    package = zip_members(source)
    names = {name.lower() for name in package.namelist()}
    try:
        xml_parts = [
            package.read(name).decode("utf-8", "replace")
            for name in package.namelist()
            if name.lower().endswith((".xml", ".rels"))
        ]
        if any(re.search(r"<!\s*(doctype|entity)\b", part, re.I) for part in xml_parts):
            package.close()
            refuse("DTD and entity declarations in Office packages are unsupported")
        content_types = package.read("[Content_Types].xml").decode("utf-8", "replace").lower()
        relationships = "\n".join(
            package.read(name).decode("utf-8", "replace").lower()
            for name in package.namelist()
            if name.lower().endswith(".rels")
        )
    except (KeyError, OSError, RuntimeError) as exc:
        package.close()
        refuse(f"corrupt Office package ({exc})")
    active_terms = ("macroenabled", "vbaproject", "activex", "oleobject", "externallink", "altchunk")
    has_external_relationship = re.search(r"targetmode\s*=\s*(['\"])external\1", relationships) is not None
    if has_external_relationship or any(term in content_types or term in relationships for term in active_terms):
        package.close()
        refuse("macro or active-content Office package is unsupported")
    if any(
        name.startswith(("word/embeddings/", "xl/embeddings/", "word/activex/", "xl/activex/", "customui/", "word/webextensions/", "xl/webextensions/"))
        or "/externallinks/" in name
        or name.endswith("vbaproject.bin")
        for name in names
    ):
        package.close()
        refuse("embedded or active Office content is unsupported")
    expected = "word/document.xml" if kind == "docx" else "xl/workbook.xml"
    if expected not in names:
        package.close()
        refuse(f"invalid {kind.upper()} package: required part is missing")
    return package


PDF_ACTION_KEYS = {
    "AA", "AcroForm", "EmbeddedFiles", "EmbeddedFile", "EF", "FileAttachment",
    "Filespec", "JavaScript", "JS", "Launch", "Movie", "OpenAction", "RichMedia",
    "Sound", "SubmitForm", "ResetForm", "GoToR", "GoToE", "SetOCGState", "Rendition",
    "ImportData", "Hide", "Named", "Thread", "Trans", "URI",
}
PDF_ACTION_TYPES = {
    "JavaScript", "GoTo", "Launch", "GoToR", "GoToE", "SubmitForm", "ResetForm", "ImportData",
    "Hide", "Named", "SetOCGState", "Rendition", "Movie", "Sound", "URI",
}


def pdf_tokens(source: str) -> list[tuple[str, str]]:
    """Tokenize PDF object syntax, decoding #xx escapes in names, not strings/comments."""
    tokens = []
    index = 0
    delimiters = "()<>[]{}/%"
    while index < len(source):
        char = source[index]
        if char.isspace():
            index += 1
            continue
        if char == "%":
            newline = source.find("\n", index)
            index = len(source) if newline < 0 else newline + 1
            continue
        if char == "(":
            depth = 1
            index += 1
            while index < len(source) and depth:
                if source[index] == "\\":
                    index += 2
                elif source[index] == "(":
                    depth += 1
                    index += 1
                elif source[index] == ")":
                    depth -= 1
                    index += 1
                else:
                    index += 1
            tokens.append(("string", ""))
            continue
        if source.startswith("<<", index):
            tokens.append(("dict-start", ""))
            index += 2
            continue
        if source.startswith(">>", index):
            tokens.append(("dict-end", ""))
            index += 2
            continue
        if char == "<":
            end = source.find(">", index + 1)
            index = len(source) if end < 0 else end + 1
            tokens.append(("string", ""))
            continue
        if char == "/":
            end = index + 1
            while end < len(source) and not source[end].isspace() and source[end] not in delimiters:
                end += 1
            raw_name = source[index + 1:end]
            name = re.sub(r"#([0-9A-Fa-f]{2})", lambda match: chr(int(match.group(1), 16)), raw_name)
            tokens.append(("name", name))
            index = end
            continue
        if char in "[]{}":
            tokens.append((char, char))
            index += 1
            continue
        end = index + 1
        while end < len(source) and not source[end].isspace() and source[end] not in delimiters:
            end += 1
        tokens.append(("word", source[index:end]))
        index = end
    return tokens


def pdf_object_has_action(object_text: str) -> bool:
    tokens = pdf_tokens(object_text)

    def value(index: int) -> tuple[tuple[str, str] | None, int]:
        if index >= len(tokens):
            return None, index
        kind, name = tokens[index]
        if kind == "dict-start":
            has_action, end = dictionary(index + 1)
            return ("action-dict", "") if has_action else ("dict", ""), end
        if kind == "[":
            index += 1
            while index < len(tokens) and tokens[index][0] != "]":
                _, index = value(index)
            return ("array", ""), index + 1
        return (kind, name), index + 1

    def dictionary(index: int) -> tuple[bool, int]:
        while index < len(tokens) and tokens[index][0] != "dict-end":
            key_kind, key = tokens[index]
            if key_kind != "name":
                _, index = value(index)
                continue
            item, index = value(index + 1)
            if key in PDF_ACTION_KEYS:
                return True, index
            if key == "S" and item and item[0] == "name" and item[1] in PDF_ACTION_TYPES:
                return True, index
            if key == "Type" and item and item[0] == "name" and item[1] in {"Action", "Filespec", "EmbeddedFile", "RichMedia"}:
                return True, index
            if item and item[0] == "action-dict":
                return True, index
        return False, index + 1

    index = 0
    while index < len(tokens):
        if tokens[index][0] == "dict-start":
            active, index = dictionary(index + 1)
            if active:
                return True
        else:
            index += 1
    return False


def extract_pdf(source: Path) -> dict:
    try:
        import fitz

        document = fitz.open(source)
        if document.is_encrypted:
            refuse("encrypted PDF is unsupported")
        if document.xref_length() > 100_000:
            refuse("PDF object count exceeds 100,000-object limit")
        # Every parsed indirect object is inspected, including page, annotation, name-tree,
        # object-stream and action dictionaries. PDF names are decoded before key matching.
        for xref in range(1, document.xref_length()):
            try:
                object_text = document.xref_object(xref, compressed=False)
            except Exception:
                refuse("PDF contains an unreadable object structure")
            if pdf_object_has_action(object_text):
                refuse("PDF contains an unsupported action, attachment, link, or interactive structure")
        if pdf_object_has_action(document.pdf_trailer()):
            refuse("PDF trailer contains an unsupported action or interactive structure")
        if document.page_count > MAX_PDF_PAGES:
            refuse(f"PDF exceeds {MAX_PDF_PAGES}-page limit")
        pages = []
        for number, page in enumerate(document, 1):
            text = page.get_text("text").strip()
            if not text:
                refuse(f"PDF page {number} has no text layer (scan/OCR required)")
            pages.append({"page": number, "text": text})
        if not pages:
            refuse("PDF contains no pages")
        return {"format": "pdf", "pages": pages}
    except Refusal:
        raise
    except Exception as exc:
        refuse(f"unreadable or corrupt PDF ({exc})")


DOCX_STORY_PART = re.compile(r"^word/(?:header\d+|footer\d+|footnotes|endnotes|comments)\.xml$", re.I)
DOCX_SUPPORTED_PARTS = {"word/document.xml"}
DOCX_WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
DOCX_NS = {"w": DOCX_WORD_NS}


def xml_text(element, exclude_textboxes: bool = False) -> str:
    text_tags = {f"{{{DOCX_WORD_NS}}}t", f"{{{DOCX_WORD_NS}}}delText"}
    values = []
    for node in element.iter():
        if node.tag not in text_tags and (not isinstance(node.tag, str) or node.tag.rsplit("}", 1)[-1] != "t"):
            continue
        parent = node.getparent()
        in_textbox = False
        while exclude_textboxes and parent is not None and parent is not element:
            if parent.tag == f"{{{DOCX_WORD_NS}}}txbxContent":
                in_textbox = True
                break
            parent = parent.getparent()
        if not in_textbox:
            values.append(node.text or "")
    return "".join(values).strip()


def extract_docx_auxiliary_part(part_name: str, payload: bytes) -> list[dict]:
    from lxml import etree

    root = etree.fromstring(payload, parser=etree.XMLParser(resolve_entities=False, no_network=True))
    for revision in ("ins", "del", "moveFrom", "moveTo"):
        if root.xpath(f".//w:{revision}[.//w:t or .//w:delText]", namespaces=DOCX_NS):
            refuse(f"DOCX tracked-change text is unsupported in {part_name}")
    blocks = []
    story_prefix = part_name
    paragraphs = root.xpath(".//w:p[not(ancestor::w:tbl)][not(ancestor::w:txbxContent)]", namespaces=DOCX_NS)
    for index, paragraph in enumerate(paragraphs, 1):
        text = xml_text(paragraph, exclude_textboxes=True)
        if text:
            blocks.append({"locator": f"{story_prefix} / paragraph {index}", "text": text})
    for table_index, table in enumerate(root.xpath(".//w:tbl[not(ancestor::w:txbxContent)]", namespaces=DOCX_NS), 1):
        if table.xpath("ancestor::w:tc", namespaces=DOCX_NS):
            refuse(f"DOCX nested table is unsupported in {part_name}")
        for row_index, row in enumerate(table.xpath("./w:tr", namespaces=DOCX_NS), 1):
            for col_index, cell in enumerate(row.xpath("./w:tc", namespaces=DOCX_NS), 1):
                text = xml_text(cell, exclude_textboxes=True)
                if text:
                    blocks.append({"locator": f"{story_prefix} / table {table_index} row {row_index} cell {col_index}", "text": text})
    for box_index, textbox in enumerate(root.xpath(".//w:txbxContent", namespaces=DOCX_NS), 1):
        for paragraph_index, paragraph in enumerate(textbox.xpath(".//w:p[not(ancestor::w:tbl)]", namespaces=DOCX_NS), 1):
            text = xml_text(paragraph)
            if text:
                blocks.append({"locator": f"{story_prefix} / text box {box_index} / paragraph {paragraph_index}", "text": text})
        for table_index, table in enumerate(textbox.xpath(".//w:tbl", namespaces=DOCX_NS), 1):
            for row_index, row in enumerate(table.xpath("./w:tr", namespaces=DOCX_NS), 1):
                for col_index, cell in enumerate(row.xpath("./w:tc", namespaces=DOCX_NS), 1):
                    text = xml_text(cell)
                    if text:
                        blocks.append({"locator": f"{story_prefix} / text box {box_index} / table {table_index} row {row_index} cell {col_index}", "text": text})
    return blocks


PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_.]*)\s*\}\}")
TAG = re.compile(r"\{%\s*(.*?)\s*%\}")


def docx_structure(blocks: list[dict], main_root, names: set[str]) -> dict:
    """Placeholders, loop tags, tables, images and styles of a DOCX reference.

    Reports what the text says; it proposes nothing. Unrecognised braces are
    listed so a reviewer sees them before any render."""
    placeholders: list[str] = []
    table_placeholders: list[str] = []
    loops: list[dict] = []
    unrecognized: list[dict] = []
    for block in blocks:
        text = block["text"]
        found = PLACEHOLDER.findall(text)
        for name in found:
            if name not in placeholders:
                placeholders.append(name)
            if "table " in block["locator"] and name not in table_placeholders:
                table_placeholders.append(name)
        for tag in TAG.findall(text):
            if tag.startswith(("for ", "endfor", "if ", "endif")):
                loops.append({"locator": block["locator"], "tag": tag})
            else:
                unrecognized.append({"locator": block["locator"], "text": text[:200]})
        leftover = PLACEHOLDER.sub("", TAG.sub("", text))
        if "{" in leftover or "}" in leftover:
            unrecognized.append({"locator": block["locator"], "text": text[:200]})
    tables = main_root.xpath("//w:tbl[not(ancestor::w:txbxContent)]", namespaces=DOCX_NS)
    images = main_root.xpath("//w:drawing | //w:pict", namespaces=DOCX_NS)
    return {
        "text_blocks": len(blocks),
        "placeholders": placeholders,
        "table_placeholders": table_placeholders,
        "loops": loops,
        "unrecognized_markers": unrecognized,
        "tables": len(tables),
        "images": len(images),
        "media_parts": sorted(name for name in names if name.startswith("word/media/")),
        "styles": sorted({block["style"] for block in blocks if block.get("style")}),
    }


def extract_docx(source: Path) -> dict:
    package = preflight_office(source, "docx")
    try:
        from docx import Document
        from lxml import etree

        names = set(package.namelist())
        supported_parts = {name for name in names if DOCX_STORY_PART.fullmatch(name)} | DOCX_SUPPORTED_PARTS
        for name in names:
            if name.lower().startswith("word/") and name.lower().endswith(".xml") and name not in supported_parts:
                root = etree.fromstring(package.read(name), parser=etree.XMLParser(resolve_entities=False, no_network=True))
                meaningful = root.xpath("//*[local-name()='t' or local-name()='delText'][normalize-space(text())]")
                if meaningful:
                    refuse(f"DOCX contains meaningful text in unsupported part {name}")

        document = Document(source)
        main_root = etree.fromstring(package.read("word/document.xml"), parser=etree.XMLParser(resolve_entities=False, no_network=True))
        body = main_root.find(".//w:body", DOCX_NS)
        for child in body:
            if child.tag not in {f"{{{DOCX_WORD_NS}}}p", f"{{{DOCX_WORD_NS}}}tbl", f"{{{DOCX_WORD_NS}}}sectPr"} and xml_text(child):
                refuse("DOCX contains meaningful body text in an unsupported structure")
        for revision in ("ins", "del", "moveFrom", "moveTo"):
            if main_root.xpath(f".//w:{revision}[.//w:t or .//w:delText]", namespaces=DOCX_NS):
                refuse("DOCX tracked-change text is unsupported")
        for part_name in supported_parts:
            root = main_root if part_name == "word/document.xml" else etree.fromstring(package.read(part_name), parser=etree.XMLParser(resolve_entities=False, no_network=True))
            drawings = root.xpath("//*[local-name()='docPr' or local-name()='cNvPr']")
            if any(value.strip() for element in drawings for key, value in element.attrib.items() if key.rsplit("}", 1)[-1] in {"descr", "title"}):
                refuse(f"DOCX has unsupported drawing description text in {part_name}")
        blocks = []
        heading = "document start"
        for index, paragraph in enumerate(document.paragraphs, 1):
            text = xml_text(paragraph._p, exclude_textboxes=True)
            style = paragraph.style.name if paragraph.style else ""
            if style.lower().startswith("heading") and text:
                heading = text
            if text:
                blocks.append({"locator": f"word/document.xml / {heading} / paragraph {index}", "style": style, "text": text})
        for table_index, table in enumerate(document.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                for col_index, cell in enumerate(row.cells, 1):
                    text = xml_text(cell._tc, exclude_textboxes=True)
                    if text:
                        blocks.append({"locator": f"word/document.xml / table {table_index} row {row_index} cell {col_index}", "text": text})
        for part_name in sorted(supported_parts - DOCX_SUPPORTED_PARTS):
            blocks.extend(extract_docx_auxiliary_part(part_name, package.read(part_name)))
        main_xml = main_root
        blocks.extend(
            block for block in extract_docx_auxiliary_part("word/document.xml", etree.tostring(main_xml))
            if "text box" in block["locator"]
        )
        if not blocks:
            refuse("DOCX contains no readable paragraph or table text")
        return {"format": "docx", "blocks": blocks, "structure": docx_structure(blocks, main_root, names)}
    except Refusal:
        raise
    except Exception as exc:
        refuse(f"unreadable or corrupt DOCX ({exc})")
    finally:
        package.close()


def extract_xlsx(source: Path) -> dict:
    package = preflight_office(source, "xlsx")
    package.close()
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        sheets = []
        count = 0
        for sheet in workbook.worksheets:
            rows = []
            for row_number, row in enumerate(sheet.iter_rows(), 1):
                values = []
                for col_number, cell in enumerate(row, 1):
                    if cell.data_type == "f":
                        refuse(f"XLSX formula in {sheet.title}!{cell.coordinate} is not a verified input value")
                    value = cell.value
                    if value is not None:
                        count += 1
                        if count > MAX_XLSX_CELLS:
                            refuse(f"XLSX exceeds {MAX_XLSX_CELLS}-nonempty-cell limit")
                        record = {"cell": cell.coordinate, "value": value}
                        if isinstance(value, datetime.datetime):
                            record.update(value=value.isoformat(), value_type="datetime")
                        elif isinstance(value, datetime.date):
                            record.update(value=value.isoformat(), value_type="date")
                        elif isinstance(value, datetime.time):
                            record.update(value=value.isoformat(), value_type="time")
                        elif isinstance(value, float) and not math.isfinite(value):
                            refuse(f"unsupported non-finite XLSX value in {sheet.title}!{cell.coordinate}")
                        elif not isinstance(value, (str, int, float, bool)):
                            refuse(f"unsupported XLSX value type {type(value).__name__} in {sheet.title}!{cell.coordinate}")
                        values.append(record)
                if values:
                    rows.append({"row": row_number, "cells": values})
            sheets.append({"sheet": sheet.title, "rows": rows})
        workbook.close()
        if count == 0:
            refuse("XLSX contains no non-empty values")
        return {"format": "xlsx", "sheets": sheets}
    except Refusal:
        raise
    except Exception as exc:
        refuse(f"unreadable or corrupt XLSX ({exc})")


def safe_output_dir(path: Path, source: Path) -> Path:
    target = path.expanduser().absolute()
    if target.is_symlink() or any(parent.is_symlink() for parent in target.parents):
        refuse("output directory path must not contain symbolic links")
    try:
        target.mkdir(parents=True, exist_ok=True)
        resolved = target.resolve(strict=True)
    except (OSError, RuntimeError) as exc:
        refuse(f"cannot prepare output directory ({exc})")
    if not resolved.is_dir() or resolved == source.parent or resolved in source.parents:
        refuse("output directory must not contain the source or be its parent")
    if source.parent == resolved:
        refuse("output directory must differ from source directory")
    return resolved


def parser_identity(kind: str) -> dict:
    from importlib.metadata import version

    parsers = {
        "pdf": {"name": "PyMuPDF", "version": version("PyMuPDF")},
        "docx": {"name": "python-docx", "version": version("python-docx")},
        "xlsx": {"name": "openpyxl", "version": version("openpyxl")},
    }
    return parsers[kind]


def create_output_file(directory_fd: int, name: str, payload: bytes | None = None) -> int:
    flags = os.O_WRONLY | os.O_CREAT | os.O_EXCL | getattr(os, "O_NOFOLLOW", 0)
    fd = os.open(name, flags, 0o600, dir_fd=directory_fd)
    if payload is not None:
        try:
            with os.fdopen(fd, "wb", closefd=False) as stream:
                stream.write(payload)
                stream.flush()
                os.fsync(fd)
        except BaseException:
            os.close(fd)
            os.unlink(name, dir_fd=directory_fd)
            raise
    return fd


PDF_LIMITATIONS = [
    "Reconstructed from the PDF text layer only: table grids, fonts, logos, colours and exact positions are not reproduced.",
    "Text order follows the PDF text stream and may not match the visual reading order.",
    "This DOCX is a reference to rebuild a template from. It is not an approved template and needs its own review and approval.",
]


def pdf_reconstruction(extracted: dict, source_name: str) -> bytes:
    """An editable DOCX holding the PDF's text per page, with its limitations."""
    from docx import Document

    document = Document()
    document.add_heading("Reconstructed reference (not an approved template)", level=1)
    document.add_paragraph(f"Source: {source_name}")
    for limitation in PDF_LIMITATIONS:
        document.add_paragraph(limitation)
    for page in extracted["pages"]:
        document.add_heading(f"PDF page {page['page']}", level=2)
        for line in page["text"].splitlines():
            if line.strip():
                document.add_paragraph(line.strip())
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def process_snapshot(snapshot: dict, output_dir: Path) -> None:
    source = snapshot["resolved"]
    snapshot_path = snapshot["snapshot"]
    source_name = snapshot["filename"]
    output = safe_output_dir(output_dir, source)
    if snapshot_path.suffix == ".pdf":
        extracted = extract_pdf(snapshot_path)
    elif snapshot_path.suffix == ".docx":
        extracted = extract_docx(snapshot_path)
    else:
        extracted = extract_xlsx(snapshot_path)
    verify_source_snapshot(snapshot)

    original_name = "source" + snapshot_path.suffix
    provenance_name = f"{Path(source_name).stem}.provenance.json"
    original = output / original_name
    provenance = output / provenance_name
    out_fd = os.open(output, os.O_RDONLY | getattr(os, "O_DIRECTORY", 0) | getattr(os, "O_NOFOLLOW", 0))
    created: list[str] = []
    try:
        original_fd = create_output_file(out_fd, original_name)
        created.append(original_name)
        preserved_hash = hashlib.sha256()
        with snapshot_path.open("rb") as incoming, os.fdopen(original_fd, "wb") as copied:
            for block in iter(lambda: incoming.read(CHUNK), b""):
                preserved_hash.update(block)
                copied.write(block)
            copied.flush()
            os.fsync(copied.fileno())
        if preserved_hash.hexdigest() != snapshot["sha256"]:
            refuse("private source snapshot changed before preservation")
        verify_source_snapshot(snapshot)
        record = {
            "source": {"filename": source_name, "sha256": snapshot["sha256"], "preserved_as": original.name},
            "extractor": {"name": "business-document-intake", "version": "1.0", "python": sys.version.split()[0], "parser": parser_identity(extracted["format"])},
            "evidence": extracted,
        }
        if extracted["format"] == "pdf":
            reconstruction_name = f"{Path(source_name).stem}.reconstruction.docx"
            reconstruction = pdf_reconstruction(extracted, source_name)
            reconstruction_fd = create_output_file(out_fd, reconstruction_name, reconstruction)
            os.close(reconstruction_fd)
            created.append(reconstruction_name)
            record["reconstruction"] = {"file": reconstruction_name, "kind": "reference-only", "limitations": PDF_LIMITATIONS}
        payload = (json.dumps(record, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
        verify_source_snapshot(snapshot)
        provenance_fd = create_output_file(out_fd, provenance_name, payload)
        os.close(provenance_fd)
        created.append(provenance_name)
        verify_source_snapshot(snapshot)
    except BaseException:
        for name in created:
            try:
                os.unlink(name, dir_fd=out_fd)
            except FileNotFoundError:
                pass
        raise
    finally:
        os.close(out_fd)
    print(f"Extracted {source_name}; preserved original as {original}; provenance: {provenance}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    try:
        with source_snapshot(args.source) as snapshot:
            process_snapshot(snapshot, args.output_dir)
        return 0
    except Refusal as exc:
        print(str(exc), file=sys.stderr)
        return 2
    except OSError as exc:
        print(f"Input refused: file operation failed ({exc})", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
