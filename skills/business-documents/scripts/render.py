#!/usr/bin/env python3
"""Render one mapped DOCX draft and optional calculation-backed XLSX/PDF."""

from __future__ import annotations

import argparse
from decimal import Decimal, ROUND_HALF_EVEN, ROUND_HALF_UP
import hashlib
import html
import json
import math
import os
from pathlib import Path
import re
import secrets
import shutil
import subprocess
import sys
import tempfile
from typing import Any
from contextlib import ExitStack

from docxtpl import DocxTemplate, RichText
from jinja2 import Environment, StrictUndefined
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from intake import preflight_office, source_snapshot, verify_source_snapshot

MAX_XLSX_CELLS = 20_000

REQUIRED_FIELDS = {
    "project_title", "draft_date", "validity_date", "customer_name", "understanding",
    "scope", "schedule", "commercial_summary", "assumptions", "review_notes",
    "source_note", "line_items_text", "known_subtotal", "full_total",
}


def sha256(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def safe_output(directory: Path, inputs: list[Path]) -> Path:
    directory = directory.expanduser().absolute()
    if directory.is_symlink() or any(parent.is_symlink() for parent in directory.parents):
        raise ValueError("output directory path must not contain symbolic links")
    directory.mkdir(parents=True, exist_ok=True)
    target = directory.resolve(strict=True)
    if not target.is_dir():
        raise ValueError("output path is not a directory")
    for source in inputs:
        source_real = source.resolve(strict=True)
        if target == source_real.parent or target in source_real.parents:
            raise ValueError("output directory must not contain or replace an input")
    return target


def exclusive_write(path: Path, writer) -> None:
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL | getattr(os, "O_NOFOLLOW", 0), 0o600)
    try:
        with os.fdopen(fd, "wb") as stream:
            writer(stream)
    except BaseException:
        path.unlink(missing_ok=True)
        raise


def publish_staged_file(staged: Path, output: Path, name: str) -> Path:
    """Atomically expose a completed staged file without following/replacing a target."""
    directory_fd = os.open(output, os.O_RDONLY | getattr(os, "O_DIRECTORY", 0) | getattr(os, "O_NOFOLLOW", 0))
    temporary_name = f".{name}.{secrets.token_hex(12)}.publish"
    flags = os.O_WRONLY | os.O_CREAT | os.O_EXCL | getattr(os, "O_NOFOLLOW", 0)
    temporary_fd = -1
    try:
        temporary_fd = os.open(temporary_name, flags, 0o600, dir_fd=directory_fd)
        with staged.open("rb") as source, os.fdopen(temporary_fd, "wb", closefd=False) as target:
            for block in iter(lambda: source.read(1024 * 1024), b""):
                target.write(block)
            target.flush()
            os.fsync(temporary_fd)
        os.close(temporary_fd)
        temporary_fd = -1
        os.link(temporary_name, name, src_dir_fd=directory_fd, dst_dir_fd=directory_fd, follow_symlinks=False)
        os.unlink(temporary_name, dir_fd=directory_fd)
        os.fsync(directory_fd)
        return output / name
    finally:
        if temporary_fd >= 0:
            os.close(temporary_fd)
        try:
            os.unlink(temporary_name, dir_fd=directory_fd)
        except FileNotFoundError:
            pass
        os.close(directory_fd)


def load_rates(path: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    if len(workbook.worksheets) != 1:
        raise ValueError("rate workbook must contain exactly one worksheet")
    sheet = workbook.active
    nonempty_cells = 0
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                workbook.close()
                raise ValueError(f"formula cell {sheet.title}!{cell.coordinate} is not allowed in the rate workbook")
            if cell.value is not None:
                nonempty_cells += 1
                if nonempty_cells > MAX_XLSX_CELLS:
                    workbook.close()
                    raise ValueError(f"rate workbook exceeds {MAX_XLSX_CELLS}-nonempty-cell limit")
    rows = sheet.iter_rows()
    header_cells = next(rows, None)
    headers = tuple(cell.value for cell in header_cells) if header_cells is not None else None
    core_headers = ("item_id", "description", "unit", "unit_price", "currency", "source_id", "effective_date")
    rule_headers = core_headers + ("currency_precision", "rounding_mode")
    if headers is None or tuple(headers) not in {core_headers, rule_headers}:
        raise ValueError(f"rate workbook headers must be {core_headers!r} with optional explicit rounding columns {rule_headers[7:]!r}")
    has_rounding_rules = tuple(headers) == rule_headers
    rates = {}
    for row_number, row_cells in enumerate(rows, 2):
        row = tuple(cell.value for cell in row_cells)
        if not row or all(value is None for value in row):
            continue
        item_id, description, unit, unit_price, currency, source_id, effective_date = row[:7]
        precision_value = row[7] if has_rounding_rules and len(row) > 7 else None
        rounding_value = row[8] if has_rounding_rules and len(row) > 8 else None
        if not isinstance(item_id, str):
            raise ValueError(f"item_id in rate workbook row {row_number} must be non-empty text")
        item_id = item_id.strip()
        if not item_id:
            raise ValueError(f"item_id in rate workbook row {row_number} must be non-empty text")
        if item_id in rates:
            raise ValueError(f"duplicate item_id {item_id!r} in rate workbook row {row_number}")
        if not description or not unit or not source_id or not effective_date:
            raise ValueError(f"missing description/unit/source/date in rate workbook row {row_number}")
        if not isinstance(currency, str) or not re.fullmatch(r"[A-Z]{3}", currency):
            raise ValueError(f"invalid currency in rate workbook row {row_number}")
        if isinstance(unit_price, str) and unit_price.startswith("="):
            raise ValueError(f"formula not allowed in rate workbook row {row_number}")
        if unit_price is not None and (isinstance(unit_price, bool) or not isinstance(unit_price, (int, float, Decimal)) or not math.isfinite(float(unit_price)) or unit_price < 0):
            raise ValueError(f"invalid unit_price in rate workbook row {row_number}")
        rates[item_id] = {
            "description": str(description), "unit": str(unit),
            "unit_price": Decimal(str(unit_price)) if unit_price is not None else None,
            "currency": str(currency), "source_id": str(source_id),
            "effective_date": str(effective_date),
            "currency_precision": precision_value if isinstance(precision_value, int) and not isinstance(precision_value, bool) and 0 <= precision_value <= 4 else None,
            "rounding_mode": str(rounding_value) if rounding_value is not None else None,
        }
    workbook.close()
    return rates


ROUNDING_MODES = {"ROUND_HALF_UP": ROUND_HALF_UP, "ROUND_HALF_EVEN": ROUND_HALF_EVEN}


def money_text(currency: str, amount: Decimal | None, precision: int | None) -> str:
    if amount is None or precision is None:
        return "UNRESOLVED"
    return f"{currency} {amount:,.{precision}f}"


def calculate(brief: dict, rates: dict[str, dict[str, Any]]) -> dict:
    lines = brief.get("price_items", [])
    selected_rates = [rates[line["item_id"]] for line in lines if line["item_id"] in rates and rates[line["item_id"]]["unit_price"] is not None]
    rule_pairs = {(rate["currency_precision"], rate["rounding_mode"]) for rate in selected_rates}
    common_rule = next(iter(rule_pairs)) if len(rule_pairs) == 1 else None
    if common_rule and (common_rule[0] is None or common_rule[1] not in ROUNDING_MODES):
        common_rule = None
    items = []
    known_total = Decimal("0")
    unresolved = False
    rounding_unresolved = bool(selected_rates and common_rule is None)
    for line in lines:
        item_id = line["item_id"]
        if item_id not in rates:
            raise ValueError(f"no approved rate row for {item_id}")
        rate = rates[item_id]
        quantity = Decimal(str(line["quantity"]))
        if not quantity.is_finite() or quantity < 0:
            raise ValueError(f"negative quantity for {item_id}")
        if rate["currency"] != str(brief.get("currency", "")):
            raise ValueError(f"currency mismatch for {item_id}")
        unit_price = rate["unit_price"]
        precision = rate["currency_precision"]
        mode_name = rate["rounding_mode"]
        supported_rule = common_rule is not None and (precision, mode_name) == common_rule
        missing = unit_price is None
        amount = None
        if not missing and supported_rule:
            quantum = Decimal(1).scaleb(-precision)
            amount = (unit_price * quantity).quantize(quantum, rounding=ROUNDING_MODES[mode_name])
        elif not missing:
            rounding_unresolved = True
        unresolved |= (missing or (not supported_rule and not missing)) and bool(line.get("mandatory", True))
        if amount is not None:
            known_total += amount
        formula = (f"=ROUND({unit_price}*{quantity},{precision},{mode_name})" if amount is not None
                   else "=UNRESOLVED (approved rate missing)" if missing
                   else "=UNRESOLVED (missing or unsupported currency rounding rule)")
        items.append({
            "description": rate["description"], "quantity": str(quantity), "unit": rate["unit"],
            "unit_price": None if missing else float(unit_price), "currency": rate["currency"],
            "currency_precision": precision, "rounding_mode": mode_name,
            "amount": None if amount is None else float(amount), "amount_text": money_text(rate["currency"], amount, precision),
            "formula_text": formula, "source_id": rate["source_id"],
            "effective_date": rate["effective_date"], "input_citation": line["input_citation"],
            "mandatory": bool(line.get("mandatory", True)),
        })
    precision = common_rule[0] if common_rule else None
    if precision is not None:
        quantum = Decimal(1).scaleb(-precision)
        known_total = known_total.quantize(quantum, rounding=ROUNDING_MODES[common_rule[1]])
    known_subtotal = None if rounding_unresolved or precision is None else float(known_total)
    return {
        "items": items,
        "known_subtotal": known_subtotal,
        "currency": str(brief.get("currency", "")),
        "currency_precision": precision,
        "rounding_mode": common_rule[1] if common_rule else None,
        "full_total": None if unresolved or rounding_unresolved else float(known_total),
        "unresolved_mandatory": unresolved or rounding_unresolved,
    }


def format_currency(currency: str, amount: float | None, precision: int | None) -> str:
    if amount is None or precision is None:
        return "UNRESOLVED"
    return f"{currency} {amount:,.{precision}f}"


def write_xlsx(path: Path, result: dict) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Suggested quotation"
    headers = ["Description", "Quantity", "Unit", "Unit price", "Amount", "Formula (text)", "Rate source/date", "Input citation"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="204A67")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for item in result["items"]:
        row = [item["description"], float(item["quantity"]), item["unit"], item["unit_price"] if item["unit_price"] is not None else "UNRESOLVED", item["amount"] if item["amount"] is not None else "UNRESOLVED", item["formula_text"], f"{item['source_id']} ({item['effective_date']})", item["input_citation"]]
        sheet.append(row)
        for text_column in (1, 3, 6, 7, 8):
            sheet.cell(sheet.max_row, text_column).data_type = "s"
        formula_cell = sheet.cell(sheet.max_row, 6)
        formula_cell.number_format = "@"
        precision = item["currency_precision"]
        number_format = '"' + result["currency"] + ' "#,##0' + ("." + "0" * precision if precision is not None and precision else "")
        sheet.cell(sheet.max_row, 4).number_format = number_format if item["unit_price"] is not None else "General"
        sheet.cell(sheet.max_row, 5).number_format = number_format if item["amount"] is not None else "General"
    sheet.append(["Known-priced subtotal", None, None, None, result["known_subtotal"] if result["known_subtotal"] is not None else "UNRESOLVED", "sum of known line amounts", "", "Excludes unresolved mandatory prices"])
    precision = result["currency_precision"]
    sheet.cell(sheet.max_row, 5).number_format = ('"' + result["currency"] + ' "#,##0' + ("." + "0" * precision if precision else "")) if precision is not None and result["known_subtotal"] is not None else "General"
    sheet.append(["Full quotation total", None, None, None, result["full_total"] if result["full_total"] is not None else "UNRESOLVED", "UNRESOLVED when a mandatory rate or rounding rule is missing", "", "Not an approved quotation"])
    if precision is not None and result["full_total"] is not None:
        sheet.cell(sheet.max_row, 5).number_format = '"' + result["currency"] + ' "#,##0' + ("." + "0" * precision if precision else "")
    sheet.append(["Status", "SUGGESTED — NOT APPROVED — HUMAN REVIEW REQUIRED"])
    sheet.freeze_panes = "A2"
    widths = [38, 13, 18, 18, 18, 44, 34, 38]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.auto_filter.ref = f"A1:H{1 + len(result['items'])}"
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False
    workbook.calculation.calcMode = "manual"
    exclusive_write(path, workbook.save)


def check_template(template: Path) -> None:
    from docx import Document

    package = preflight_office(template, "docx")
    try:
        xml_parts = [package.read(name).decode("utf-8", "replace") for name in package.namelist() if name.lower().endswith((".xml", ".rels"))]
    finally:
        package.close()
    if any(re.search(r"<!\s*(doctype|entity)\b", part, re.I) for part in xml_parts):
        raise ValueError("prepared template contains an unsupported DTD/entity declaration")
    combined_text = html.unescape("\n".join(re.sub(r"<[^>]*>", "", part) for part in xml_parts))
    expressions = re.findall(r"\{\{.*?\}\}|\{%.*?%\}", combined_text, re.S)
    if combined_text.count("{{") != sum(token.startswith("{{") for token in expressions) or combined_text.count("{%") != sum(token.startswith("{%") for token in expressions):
        raise ValueError("template contains an unbalanced template expression")
    allowed_fields = REQUIRED_FIELDS | {"is_financial", "line_items"}
    for token in expressions:
        if token.startswith("{{"):
            expression = token[2:-2].strip()
            if expression.startswith("r "):
                expression = expression[2:].strip()
            if expression not in allowed_fields:
                raise ValueError(f"prepared template uses an unmapped expression: {expression[:80]}")
        elif token[2:-2].strip() not in {"p if is_financial", "p endif"}:
            raise ValueError(f"prepared template uses an unsupported control tag: {token[:80]}")
    document = Document(template)
    text = "\n".join([p.text for p in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    missing = [field for field in REQUIRED_FIELDS if ("{{r " + field + " }}" if field == "line_items_text" else "{{ " + field + " }}") not in text]
    if missing:
        raise ValueError("prepared template missing mapped required fields: " + ", ".join(sorted(missing)))
    if not document.sections or not any(section.header.paragraphs for section in document.sections) or not any(section.footer.paragraphs for section in document.sections):
        raise ValueError("prepared template must include a header and footer")


def check_task_input(path: Path, suffix: str) -> None:
    if path.suffix.lower() != suffix or path.is_symlink() or not path.is_file():
        raise ValueError(f"input must be a regular non-symlink {suffix} file")
    if path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("input exceeds 20 MiB limit")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--brief", required=True, type=Path)
    parser.add_argument("--rates", type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    output = None
    created_paths: list[Path] = []
    try:
        with ExitStack() as snapshots_stack:
            snapshots = {
                "template": snapshots_stack.enter_context(source_snapshot(args.template, {".docx"})),
                "brief": snapshots_stack.enter_context(source_snapshot(args.brief, {".json"})),
            }
            if args.rates:
                snapshots["rates"] = snapshots_stack.enter_context(source_snapshot(args.rates, {".xlsx"}))
            template_input = snapshots["template"]["snapshot"]
            brief_input = snapshots["brief"]["snapshot"]
            rates_input = snapshots["rates"]["snapshot"] if "rates" in snapshots else None
            check_task_input(template_input, ".docx")
            check_task_input(brief_input, ".json")
            if rates_input:
                check_task_input(rates_input, ".xlsx")
            for snapshot in snapshots.values():
                verify_source_snapshot(snapshot)
            inputs = [snapshot["resolved"] for snapshot in snapshots.values()]
            output = safe_output(args.output_dir, inputs)
            brief = json.loads(brief_input.read_text(encoding="utf-8"))
            if not isinstance(brief, dict) or not isinstance(brief.get("price_items", []), list):
                raise ValueError("brief must be a JSON object with a price_items list")
            financial = bool(brief.get("financial", False))
            if financial and rates_input is None:
                raise ValueError("financial brief requires an approved rate workbook")
            if financial and not brief.get("price_items"):
                raise ValueError("financial brief requires at least one explicitly sourced price item")
            if not financial and brief.get("price_items"):
                raise ValueError("non-financial brief must not contain price_items")
            if financial:
                rates_package = preflight_office(rates_input, "xlsx")
                rates_package.close()
            if financial:
                rates = load_rates(rates_input)
                for row_number, line in enumerate(brief["price_items"], 1):
                    if not isinstance(line, dict) or not isinstance(line.get("item_id"), str):
                        raise ValueError(f"price_items entry {row_number} item_id must be non-empty text")
                    item_id = line["item_id"].strip()
                    if not item_id:
                        raise ValueError(f"price_items entry {row_number} item_id must be non-empty text")
                    line["item_id"] = item_id
                    if item_id not in rates:
                        raise ValueError(f"no approved rate row for item_id {item_id!r}")
                result = calculate(brief, rates)
            else:
                result = {
                    "items": [], "known_subtotal": None, "currency": "", "full_total": None,
                    "unresolved_mandatory": False,
                }
            context = dict(brief["template_fields"])
            if financial:
                item_lines = []
                for item in result["items"]:
                    unit_price = "UNRESOLVED" if item["unit_price"] is None else f"{item['currency']} {item['unit_price']:,} per {item['unit']}"
                    item_lines.append(f"{item['description']}: quantity {item['quantity']} {item['unit']}; {unit_price}; extension {item['amount_text']}; rate {item['source_id']} effective {item['effective_date']}; input {item['input_citation']}.")
                line_items_text = RichText()
                for index, line in enumerate(item_lines):
                    line_items_text.add(line)
                    if index + 1 < len(item_lines):
                        line_items_text.xml += "<w:br/>"
            else:
                line_items_text = RichText("Not applicable — no tender pricing or quotation requested.")
            context.update({
                "is_financial": financial,
                "line_items": result["items"],
                "line_items_text": line_items_text,
                "known_subtotal": format_currency(result["currency"], result["known_subtotal"], result["currency_precision"]) if financial else "Not applicable — non-financial task",
                "full_total": format_currency(result["currency"], result["full_total"], result["currency_precision"]) if financial else "Not applicable — non-financial task",
            })
            missing = REQUIRED_FIELDS - context.keys()
            if missing:
                raise ValueError("brief lacks mapped template fields: " + ", ".join(sorted(missing)))
            check_template(template_input)
            stem = str(brief.get("output_stem", "business-draft"))
            if not stem or any(char not in "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789-_" for char in stem):
                raise ValueError("output_stem must contain only letters, digits, hyphens or underscores")
            docx_name = f"{stem}.docx"
            pdf_name = f"{stem}.pdf"
            xlsx_name = f"{stem}-quotation.xlsx"
            metadata_name = f"{stem}.render.json"
            reserved_names = [docx_name, pdf_name, metadata_name] + ([xlsx_name] if financial else [])
            if any(os.path.lexists(output / name) for name in reserved_names):
                raise ValueError("an output already exists (including symlinks); refusing to overwrite")
            import fitz

            with tempfile.TemporaryDirectory(prefix="business-render-", dir="/tmp") as staging_dir:
                staging = Path(staging_dir)
                staged_docx = staging / docx_name
                staged_xlsx = staging / xlsx_name
                staged_pdf = staging / pdf_name
                staged_metadata = staging / metadata_name
                template = DocxTemplate(str(template_input))
                template.render(context, jinja_env=Environment(undefined=StrictUndefined, autoescape=True))
                exclusive_write(staged_docx, template.save)
                if financial:
                    write_xlsx(staged_xlsx, result)
                with tempfile.TemporaryDirectory(prefix="cadence-lo-", dir=staging_dir) as converter_dir:
                    converted_pdf = Path(converter_dir) / pdf_name
                    command = ["soffice", "--headless", f"-env:UserInstallation=file://{converter_dir}/profile", "--convert-to", "pdf", "--outdir", converter_dir, str(staged_docx)]
                    subprocess.run(command, check=True, capture_output=True, text=True, timeout=85)
                    if converted_pdf.is_symlink() or not converted_pdf.is_file() or converted_pdf.stat().st_size == 0:
                        raise ValueError("LibreOffice did not create the expected PDF in private staging")
                    shutil.copyfile(converted_pdf, staged_pdf)
                pdf = fitz.open(staged_pdf)
                if pdf.page_count < 1 or pdf.page_count > 100:
                    raise ValueError("rendered PDF page count is outside supported bounds")
                page_images = []
                total_preview_bytes = 0
                for page_number, page in enumerate(pdf, 1):
                    if not page.get_text("text").strip():
                        raise ValueError("rendered PDF contains a page without extractable text")
                    image_bytes = page.get_pixmap(matrix=fitz.Matrix(1.0, 1.0), alpha=False).tobytes("png")
                    total_preview_bytes += len(image_bytes)
                    if total_preview_bytes > 64 * 1024 * 1024:
                        raise ValueError("PDF page previews exceed 64 MiB output bound")
                    image_name = f"{stem}-page-{page_number:03d}.png"
                    exclusive_write(staging / image_name, lambda stream, payload=image_bytes: stream.write(payload))
                    page_images.append(image_name)
                pdf.close()
                meta = {
                    "draft": {"docx": docx_name, "pdf_from_docx": pdf_name, "docx_sha256": sha256(staged_docx), "pdf_sha256": sha256(staged_pdf)},
                    "template": {"filename": snapshots["template"]["filename"], "sha256": snapshots["template"]["sha256"]},
                    "brief": {"filename": snapshots["brief"]["filename"], "sha256": snapshots["brief"]["sha256"]},
                    "rates": ({"filename": snapshots["rates"]["filename"], "sha256": snapshots["rates"]["sha256"]} if financial else None),
                    "calculation": result if financial else None,
                    "page_previews": page_images,
                    "tools": {"python": sys.version.split()[0], "docxtpl": __import__("importlib.metadata", fromlist=["version"]).version("docxtpl"), "python_docx": __import__("importlib.metadata", fromlist=["version"]).version("python-docx"), "openpyxl": __import__("importlib.metadata", fromlist=["version"]).version("openpyxl"), "pymupdf": fitz.VersionBind, "libreoffice": subprocess.run(["soffice", "--version"], check=True, capture_output=True, text=True, timeout=5).stdout.strip()},
                }
                exclusive_write(staged_metadata, lambda stream: stream.write((json.dumps(meta, ensure_ascii=False, indent=2) + "\n").encode("utf-8")))
                staged_files = [(staged_docx, docx_name)]
                if financial:
                    staged_files.append((staged_xlsx, xlsx_name))
                staged_files.append((staged_pdf, pdf_name))
                staged_files.extend((staging / name, name) for name in page_images)
                staged_files.append((staged_metadata, metadata_name))
                if any(os.path.lexists(output / name) for _, name in staged_files):
                    raise ValueError("an output already exists (including symlinks); refusing to overwrite")
                for snapshot in snapshots.values():
                    verify_source_snapshot(snapshot)
                for staged, name in staged_files:
                    published = publish_staged_file(staged, output, name)
                    created_paths.append(published)
                for snapshot in snapshots.values():
                    verify_source_snapshot(snapshot)
            print(json.dumps({"status": "draft-rendered", "docx": str(output / docx_name), "pdf": str(output / pdf_name), "quotation": str(output / xlsx_name) if financial else None, "render_metadata": str(output / metadata_name), "calculation": result if financial else None}, ensure_ascii=False))
            return 0
    except Exception as exc:
        print(f"Render blocked: {type(exc).__name__}: {exc}", file=sys.stderr)
        # Remove only files published by this invocation; preserve caller files.
        for path in created_paths:
            path.unlink(missing_ok=True)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
