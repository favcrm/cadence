#!/usr/bin/env python3
"""Check one output document against its template pack (AOS-310).

The profile is the whole contract: it names the format, the fields and
their formats, the styles and sections, the tables, the input ranges that
may differ from the template, the cells whose computed values must hold,
and — optionally — a block that must stay on one printed page.

Two differences from the CAD-1345 checker this was ported from:

* formulas are recalculated by the image's LibreOffice, not the unpinned
  `formulas` package, and a missing recalculation is a FAIL (never a pass);
* `pagination.block_fits_one_page` renders the output and measures the
  printed block with PyMuPDF, so "one page wide and one page tall" is
  checked instead of judged by eye.

Structural checks only: this never edits the document, and a FAIL is a
successful run that reports a finding. Exit 0 = the check ran (whatever it
found), 2 = the input was refused before checking.

Usage:
  check.py --profile P.json --template T --output O --output-dir DIR
           [--manifest M.json]
           [--source-profile SP --source-template ST --source-output SO]
           [--expected-total "CODE 1,234.56"] [--ids REQ-001,REQ-002]
"""

from __future__ import annotations

import argparse
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import io
import json
from pathlib import Path
import re
import sys

import pack

SCHEMA = "business-documents-check@1"

FORMATS = {
    "text": r".+?",
    "token": r"\S+",
    "date": r"\d{4}-\d{2}-\d{2}",
    "date_or_sample": r"\d{4}-\d{2}-\d{2}|SAMPLE",
    "req_id": r"REQ-\d+",
    "currency": r"[A-Z]{3}",
    "money": r"[A-Z]{3} \d{1,3}(?:,\d{3})*\.\d{2}",
    "rate": r"\d+(?:\.\d+)?",
    "PAGE": r"\d*",
}
MONEY_LIKE = re.compile(r"\b[A-Z]{3} \d[\d,]*(?:\.\d+)?\b")
TOKEN = re.compile(r"\{\{|\}\}|\{%|%\}")
OPEN_MARKERS = {"[TBC]": "[TBC", "UNRESOLVED": "UNRESOLVED"}
NOT_APPLICABLE = "Not applicable"


class Report:
    def __init__(self):
        self.lines: list[str] = []
        self.fails = 0
        self.warns = 0

    def _add(self, line: str) -> None:
        self.lines.append(line)

    def ok(self, name: str, detail: str = "") -> None:
        self._add(f"PASS {name}" + (f": {detail}" if detail else ""))

    def fail(self, name: str, detail: str) -> None:
        self.fails += 1
        self._add(f"FAIL {name}: {detail}")

    def warn(self, name: str, detail: str) -> None:
        self.warns += 1
        self._add(f"WARN {name}: {detail}")

    def info(self, name: str, detail: str) -> None:
        self._add(f"INFO {name}: {detail}")

    def check(self, condition: bool, name: str, detail: str, warn_only: bool = False) -> None:
        if condition:
            self.ok(name)
        elif warn_only:
            self.warn(name, detail)
        else:
            self.fail(name, detail)

    def passes(self) -> int:
        return sum(1 for line in self.lines if line.startswith("PASS"))

    def summary(self, label: str) -> str:
        return f"SUMMARY {label}: {self.passes()} PASS, {self.warns} WARN, {self.fails} FAIL"


def fmt(profile: dict, name: str) -> str:
    return {**FORMATS, **profile.get("formats", {})}[name]


def pattern(profile: dict, template_text: str) -> re.Pattern[str]:
    """'Version {version}  |  {date}' -> regex with one named group per field."""
    out, pos = "", 0
    for match in re.finditer(r"\{(\w+)\}", template_text):
        out += re.escape(template_text[pos : match.start()])
        name = match.group(1)
        kind = profile.get("fields", {}).get(name, name if name in FORMATS else "text")
        out += f"(?P<{name}>{fmt(profile, kind)})"
        pos = match.end()
    return re.compile(out + re.escape(template_text[pos:]))


def money(code: str, value: float, places: int = 2) -> str:
    quantized = Decimal(str(value)).quantize(Decimal(1).scaleb(-places), ROUND_HALF_UP)
    return f"{code} {quantized:,.{places}f}"


def parse_money(text: str) -> tuple[str, Decimal]:
    code, amount = text.split(" ", 1)
    return code, Decimal(amount.replace(",", ""))


# ---------------------------------------------------------------- DOCX


def check_docx(profile, template_path, output_path, r, source):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    tpl, doc = Document(template_path), Document(output_path)
    ts, os_ = tpl.sections[0], doc.sections[0]
    styles = profile["styles"]

    def blocks(document):
        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, document)
            elif child.tag == qn("w:tbl"):
                yield Table(child, document)

    def cells(row):
        return [cell.text.strip() for cell in row.cells]

    def has(row, tag):
        trpr = row._tr.trPr
        return trpr is not None and trpr.find(qn(tag)) is not None

    dims = (
        "page_width",
        "page_height",
        "left_margin",
        "right_margin",
        "top_margin",
        "bottom_margin",
    )
    diff = [d for d in dims if abs((getattr(ts, d) or 0) - (getattr(os_, d) or 0)) > 3600]
    r.check(not diff, "design: page size and margins as template", f"differs: {diff}")

    tpl_styles = {style.name: style for style in tpl.styles}
    used = {p.style.name for p in doc.paragraphs} | {
        t.style.name for t in doc.tables if t.style
    }
    new = sorted(used - set(tpl_styles))
    r.check(not new, "design: only template styles used", f"styles not in template: {new}")
    changed = []
    for name in sorted(used & set(tpl_styles)):
        a, b = tpl_styles[name], doc.styles[name]
        if hasattr(a, "font") and (a.font.name, a.font.size, a.font.bold) != (
            b.font.name,
            b.font.size,
            b.font.bold,
        ):
            changed.append(name)
    r.check(not changed, "design: template style definitions unchanged", f"changed: {changed}")
    bad_tables = [
        i + 1
        for i, table in enumerate(doc.tables)
        if (table.style.name if table.style else None) != styles["table"]
    ]
    r.check(not bad_tables, f"design: tables use '{styles['table']}'", f"table(s) {bad_tables}")

    def run_fonts(document):
        found = set()
        for paragraph in document.paragraphs + [
            p
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for p in cell.paragraphs
        ]:
            for run in paragraph.runs:
                if run.font.name:
                    found.add(("font", run.font.name))
                if run.font.size:
                    found.add(("size", run.font.size.pt))
        return found

    allowed = (
        run_fonts(tpl)
        | {("font", s.font.name) for s in tpl_styles.values() if hasattr(s, "font") and s.font.name}
        | {("size", s.font.size.pt) for s in tpl_styles.values() if hasattr(s, "font") and s.font.size}
    )
    extra = sorted(run_fonts(doc) - allowed, key=str)
    r.check(
        not extra,
        "design: no direct fonts or sizes outside the template",
        f"{extra[:6]}",
        warn_only=True,
    )

    values: dict[str, set[str]] = {}

    def take(match):
        for key, value in match.groupdict().items():
            values.setdefault(key, set()).add(value.strip())

    header = " ".join(p.text for p in os_.header.paragraphs).strip()
    match = pattern(profile, profile["header"]).fullmatch(header)
    r.check(bool(match), "header", f"{header!r} does not match {profile['header']!r}")
    if match:
        take(match)
    foot = profile["footer"]
    footer = " ".join(p.text for p in os_.footer.paragraphs).strip()
    match = pattern(profile, foot["text"]).fullmatch(footer)
    r.check(bool(match), "footer", f"{footer!r} does not match {foot['text']!r}")
    if match:
        take(match)
    if foot.get("page_field"):
        xml = os_.footer._element.xml
        r.check(
            bool(re.search(r"instrText[^>]*>\s*PAGE\b|fldSimple[^>]*PAGE", xml)),
            "footer: PAGE field",
            "page number is typed text, not a field",
        )

    all_blocks = list(blocks(doc))
    heading_idx = next(
        (
            i
            for i, block in enumerate(all_blocks)
            if isinstance(block, Paragraph) and block.style.name == styles["heading"]
        ),
        len(all_blocks),
    )
    cover = [
        block.text.strip()
        for block in all_blocks[:heading_idx]
        if isinstance(block, Paragraph) and block.text.strip()
    ]
    pos = 0
    for item in profile.get("cover", []):
        rx = pattern(profile, item["text"])
        hit = next((i for i in range(pos, len(cover)) if rx.fullmatch(cover[i])), None)
        if hit is None:
            r.fail(f"cover: {item['text']}", "line missing, malformed, or out of order")
        else:
            r.ok(f"cover: {item['text']}")
            take(rx.fullmatch(cover[hit]))
            pos = hit + 1
        if item.get("toc_field"):
            r.check(
                bool(re.search(r"instrText[^>]*>\s*TOC\b|fldSimple[^>]*TOC", doc.element.xml)),
                "cover: table-of-contents field",
                "missing",
            )
    conflict = {k: sorted(v) for k, v in values.items() if len(v) > 1 and k != "PAGE"}
    r.check(not conflict, "fields consistent across header, footer and cover", f"{conflict}")
    if "SAMPLE" in values.get("date", set()):
        r.info("mode", "sample output")

    tpl_toc = {
        p.text.strip()
        for p in tpl.paragraphs
        if "TOC" in p._p.xml and "instrText" in p._p.xml and p.text.strip()
    }
    out_text = {p.text.strip() for p in doc.paragraphs}
    if tpl_toc & out_text:
        r.warn(
            "index field updated",
            "template placeholder still shown; update fields in Word before issue",
        )

    def look(run):
        color = run.font.color
        return (
            run.italic,
            run.bold,
            run.underline,
            run.font.size,
            str(color.rgb) if color is not None and color.type is not None else None,
        )

    guidance = {s["guidance"] for s in profile["sections"] if s.get("guidance")}
    hint_looks = {
        look(run)
        for p in tpl.paragraphs
        if p.text.strip() in guidance
        for run in p.runs
        if run.text.strip()
    }
    hint_looks.discard((None, None, None, None, None))
    body_paras = [b for b in all_blocks[heading_idx:] if isinstance(b, Paragraph)]
    body_paras += [
        p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs
    ]
    hinted = [
        p.text.strip()[:40]
        for p in body_paras
        if p.text.strip()
        and p.text.strip() not in guidance
        and any(look(run) in hint_looks for run in p.runs if run.text.strip())
    ]
    r.check(
        not hinted,
        "design: content not in guidance (hint) formatting",
        f"{len(hinted)} paragraph(s) look like template hints, e.g. {hinted[:2]}; write content in plain body style",
    )

    expected = [s["heading"] for s in profile["sections"]]
    found = [
        b.text.strip()
        for b in all_blocks
        if isinstance(b, Paragraph) and b.style.name == styles["heading"]
    ]
    r.check(found == expected, "sections: headings exact and in order", f"found {found}")
    content, current = {h: {"paras": [], "subs": [], "tables": []} for h in expected}, None
    for block in all_blocks[heading_idx:]:
        if isinstance(block, Paragraph) and block.style.name == styles["heading"]:
            current = block.text.strip()
        elif current in content:
            if isinstance(block, Table):
                content[current]["tables"].append(block)
            elif block.text.strip():
                content[current]["paras"].append(block.text.strip())
                if block.style.name == styles.get("subheading"):
                    content[current]["subs"].append(block.text.strip())

    for spec in profile["sections"]:
        heading, body = spec["heading"], content.get(spec["heading"])
        if body is None:
            continue
        if spec.get("guidance") and spec["guidance"] in body["paras"]:
            r.fail(f"{heading}: guidance replaced", "template guidance line still present")
        if not body["paras"] and not body["tables"]:
            r.fail(f"{heading}: has content", "section is empty (write [TBC: ...] if input is missing)")
        if spec["kind"] == "groups" and not body["subs"]:
            r.fail(f"{heading}: group subheadings", f"no '{styles.get('subheading')}' paragraphs")
        if spec["kind"] == "table":
            check_table(profile, spec, body["tables"], r, source)
        if spec.get("cite_all_ids"):
            ids = source.get("in_scope_ids")
            if ids is None:
                r.info(f"{heading}: IDs cited", "not checked (no source or ids)")
            else:
                text = " ".join(
                    body["paras"]
                    + [x for t in body["tables"] for row in t.rows for x in cells(row)]
                )
                missing = [i for i in ids if not re.search(rf"\b{re.escape(i)}\b", text)]
                r.check(not missing, f"{heading}: every in-scope ID cited", f"missing {missing}")
    if not any(line.startswith("FAIL") and "guidance" in line for line in r.lines):
        r.ok("sections: guidance lines replaced")

    text = [p.text for p in doc.paragraphs] + [
        x for table in doc.tables for row in table.rows for x in cells(row)
    ]
    text += [p.text for p in os_.header.paragraphs + os_.footer.paragraphs]
    left = [t for t in text if TOKEN.search(t)]
    r.check(not left, "no template tokens left", f"{len(left)} e.g. {left[:1]}")
    odd = sorted(
        {t for s in text for t in MONEY_LIKE.findall(s) if not re.fullmatch(fmt(profile, "money"), t)}
    )
    r.check(not odd, "money format in text", f"check {odd[:6]} (CODE 1,234.56)", warn_only=True)
    if profile.get("client_facing") and source.get("internal_refs"):
        joined = "\n".join(text)
        shown = [
            x
            for x in source["internal_refs"]
            if re.search(rf"(?<![\w-]){re.escape(x)}(?![\w-])", joined)
        ]
        r.check(
            not shown,
            "client text: no internal references",
            f"shows {shown[:6]} (cost line IDs or supplier quote refs)",
            warn_only=True,
        )

    pag = profile.get("pagination", {})
    if pag.get("heading_keep_with_next"):
        r.check(
            bool(doc.styles[styles["heading"]].paragraph_format.keep_with_next),
            "pagination: headings keep with next",
            "a heading can be stranded at a page foot",
        )
    if pag.get("repeat_header_rows"):
        bad = [i + 1 for i, table in enumerate(doc.tables) if not has(table.rows[0], "w:tblHeader")]
        r.check(not bad, "pagination: table header rows repeat", f"table(s) {bad}", warn_only=True)
    if pag.get("rows_cant_split"):
        bad = sorted(
            {i + 1 for i, table in enumerate(doc.tables) for row in table.rows if not has(row, "w:cantSplit")}
        )
        r.check(not bad, "pagination: table rows do not split", f"table(s) {bad}", warn_only=True)
    r.info(
        "open items",
        ", ".join(f"{label} x{sum(t.count(m) for t in text)}" for label, m in OPEN_MARKERS.items()),
    )


def check_table(profile, spec, tables, r, source):
    heading, tspec = spec["heading"], profile["tables"][spec["table"]]
    if not tables:
        r.fail(f"{heading}: table", "missing")
        return
    rows = [[c.text.strip() for c in row.cells] for row in tables[0].rows]
    if rows[0] != tspec["header"]:
        r.fail(f"{heading}: table header", f"{rows[0]} != {tspec['header']}")
        return
    r.ok(f"{heading}: table header")
    body = rows[1:]
    if len(body) < tspec.get("min_rows", 0):
        r.fail(f"{heading}: table rows", f"needs at least {tspec['min_rows']}")
    if tspec.get("no_empty_cells"):
        empty = [row for row in body if any(not c for c in row)]
        r.check(not empty, f"{heading}: no empty cells", f"{empty[:2]}")
    for col, kind in tspec.get("column_formats", {}).items():
        i = tspec["header"].index(col)
        bad = [
            row[i]
            for row in body
            if not re.search(fmt(profile, kind), row[i]) and not row[i].startswith("[TBC")
        ]
        r.check(not bad, f"{heading}: '{col}' format", f"{bad[:3]}")
    if "row_labels" in tspec:
        labels = [row[0] for row in body]
        r.check(labels == tspec["row_labels"], f"{heading}: rows in order", f"{labels} != {tspec['row_labels']}")
    conv = tspec.get("conversion")
    if not conv:
        return
    by = {row[0]: row for row in body}
    amount, rate_row = by.get(conv["amount_row"]), by.get(conv["rate_row"])
    if not amount or not rate_row:
        return
    local, reporting, rate = (
        amount[conv["local_column"]],
        amount[conv["reporting_column"]],
        rate_row[conv["rate_column"]],
    )
    if local == NOT_APPLICABLE:
        r.check(reporting == NOT_APPLICABLE, f"{heading}: not applicable", "local is 'Not applicable' but reporting is not")
        r.info("pricing", "not applicable")
        return
    if local == "UNRESOLVED":
        r.info("pricing", "UNRESOLVED total")
        return
    if not re.fullmatch(fmt(profile, "money"), local):
        r.fail(f"{heading}: total format", f"{local!r} is not CODE 1,234.56")
        return
    r.ok(f"{heading}: total format")
    want = source.get("local_total")
    if want:
        r.check(local == want, f"{heading}: total matches source", f"{local} != {want}")
    if reporting in (NOT_APPLICABLE, "UNRESOLVED"):
        r.info("reporting", reporting)
        return
    if not re.fullmatch(fmt(profile, "money"), reporting) or not re.fullmatch(fmt(profile, "rate"), rate):
        r.fail(f"{heading}: reporting total and rate format", f"{reporting!r}, rate {rate!r}")
        return
    _, local_amount = parse_money(local)
    _, reporting_amount = parse_money(reporting)
    places = len(local.split(".")[-1])
    calc = (local_amount * Decimal(rate)).quantize(Decimal(1).scaleb(-places), ROUND_HALF_UP)
    r.check(
        calc == reporting_amount,
        f"{heading}: reporting = local x rate",
        f"{local_amount} x {rate} = {calc}, document says {reporting_amount}",
    )
    if source.get("reporting_total"):
        r.check(
            reporting == source["reporting_total"],
            f"{heading}: reporting total matches source",
            f"{reporting} != {source['reporting_total']}",
        )
    if source.get("rate") is not None:
        r.check(
            Decimal(rate) == Decimal(str(source["rate"])),
            f"{heading}: rate matches source",
            f"{rate} != {source['rate']}",
        )


# ---------------------------------------------------------------- XLSX


def split_ref(ref: str) -> tuple[str, str]:
    sheet, cell = ref.split("!")
    return sheet, cell


def check_xlsx(profile, template_path, output_path, r, _source):
    from openpyxl import load_workbook

    tpl, wb = load_workbook(template_path), load_workbook(output_path)
    r.check(wb.sheetnames == tpl.sheetnames, "design: sheets as template, same order", f"{wb.sheetnames}")
    bounds = {
        name: pack.parse_ranges(profile.get("input_ranges", {}).get(name, []))
        for name in tpl.sheetnames
    }

    for name in tpl.sheetnames:
        if name not in wb.sheetnames:
            continue
        a, b = tpl[name], wb[name]
        changed = []
        for row in range(1, max(a.max_row, b.max_row) + 1):
            for col in range(1, max(a.max_column, b.max_column) + 1):
                if pack.inside_bounds(bounds[name], row, col):
                    continue
                if a.cell(row, col).value != b.cell(row, col).value:
                    changed.append(b.cell(row, col).coordinate)
        r.check(
            not changed,
            f"design: {name} labels and formulas unchanged outside input cells",
            f"{len(changed)} cell(s), e.g. {changed[:5]}",
        )
        validations = lambda ws: sorted(
            (str(d.sqref), d.type, d.formula1) for d in ws.data_validations.dataValidation
        )
        r.check(validations(a) == validations(b), f"design: {name} validations as template", "drop-down lists differ")
        r.check(
            a.freeze_panes == b.freeze_panes,
            f"design: {name} frozen panes as template",
            f"{b.freeze_panes}",
            warn_only=True,
        )
        if a.print_area != b.print_area:
            r.warn(f"design: {name} print area as template", f"{b.print_area}")
        if (a.page_setup.orientation, a.page_setup.fitToWidth) != (
            b.page_setup.orientation,
            b.page_setup.fitToWidth,
        ):
            r.warn(f"design: {name} page setup as template", "orientation or fit-to-width differs")

    def val(ref):
        sheet, cell = split_ref(ref)
        return wb[sheet][cell].value if sheet in wb.sheetnames else None

    missing = [c for c in profile.get("required_cells", []) if val(c) in (None, "")]
    r.check(not missing, "required cells filled", f"{missing}")
    for ref, kind in profile.get("cell_formats", {}).items():
        value = val(ref)
        if kind == "positive_number":
            ok = isinstance(value, (int, float)) and not isinstance(value, bool) and value > 0
        else:
            ok = isinstance(value, str) and re.fullmatch(fmt(profile, kind), value)
        r.check(bool(ok), f"{ref} is {kind}", f"{value!r}")

    for spec in profile.get("row_tables", []):
        check_row_table(profile, wb, spec, r)

    values = evaluate(output_path, profile, r)
    for e in profile.get("evaluate", []):
        value = values.get(e["cell"])
        r.check(
            value is not None and abs(value - e["equals"]) < 0.005,
            f"{e['cell']} = {e['equals']}",
            f"{e['reason']} (got {value})",
        )

    exports = {}
    ex = profile.get("exports", {})
    for key in ("local_total", "reporting_total"):
        spec = ex.get(key)
        if not spec:
            continue
        amount = values.get(spec["cell"])
        currency = val(spec["currency"])
        if amount is None:
            r.fail(f"export {key}", f"no numeric value at {spec['cell']} (recalculated)")
        elif not isinstance(currency, str) or not currency:
            r.fail(f"export {key}", f"no currency at {spec['currency']}")
        else:
            exports[key] = money(currency, amount)
    if ex.get("rate"):
        rate = val(ex["rate"]["cell"])
        if rate is None:
            r.fail("export rate", f"no value at {ex['rate']['cell']}")
        else:
            exports["rate"] = rate
    ids_spec = ex.get("in_scope_ids")
    if ids_spec:
        if ids_spec["sheet"] not in wb.sheetnames:
            r.fail("export in_scope_ids", f"sheet {ids_spec['sheet']!r} missing")
        else:
            ws = wb[ids_spec["sheet"]]
            ids = []
            for row in range(ids_spec["rows"][0], ids_spec["rows"][1] + 1):
                key = ws[f"{ids_spec['key']}{row}"].value
                if key and not any(
                    str(ws[f"{c}{row}"].value or "").strip() in v
                    for c, v in ids_spec.get("exclude_when", {}).items()
                ):
                    ids.append(str(key).strip())
            exports["in_scope_ids"] = ids
    refs_spec = ex.get("internal_refs")
    if refs_spec:
        if refs_spec["sheet"] not in wb.sheetnames:
            r.fail("export internal_refs", f"sheet {refs_spec['sheet']!r} missing")
        else:
            ws, refs = wb[refs_spec["sheet"]], set()
            for row in range(refs_spec["rows"][0], refs_spec["rows"][1] + 1):
                for col in refs_spec["columns"]:
                    first = (
                        str(ws[f"{col}{row}"].value or "")
                        .split(refs_spec.get("split", "\n"))[0]
                        .split()
                    )
                    if first and not first[0].startswith("[TBC"):
                        refs.add(first[0])
            exports["internal_refs"] = sorted(refs)
    for key, value in exports.items():
        r.info(f"export {key}", ", ".join(value) if isinstance(value, list) else value)
    return exports


def check_row_table(profile, wb, spec, r):
    name = spec["sheet"]
    if name not in wb.sheetnames:
        r.fail(f"{name}: sheet present", "the output has no such sheet")
        return
    ws = wb[name]
    first, last = spec["rows"]

    def get(col, row):
        return ws[f"{col}{row}"].value

    def lookup(ref):
        return {
            str(cell.value).strip()
            for row in wb[split_ref(ref)[0]][split_ref(ref)[1]]
            for cell in row
            if cell.value not in (None, "")
        }

    keys, problems, warns = [], [], []
    for row in range(first, last + 1):
        key = get(spec["key"], row)
        if key in (None, ""):
            continue
        key = str(key).strip()
        keys.append(key)
        if spec.get("key_format") and not re.fullmatch(fmt(profile, spec["key_format"]), key):
            problems.append(f"{key}: ID format")
        for col in spec.get("required", []):
            if get(col, row) in (None, ""):
                problems.append(f"{key}: {col} empty")
        for rule in spec.get("rules", []):
            if all(str(get(c, row) or "").strip() in v for c, v in rule["when"].items()):
                for col in rule["required"]:
                    if get(col, row) in (None, ""):
                        problems.append(f"{key}: {col} empty ({rule['reason']})")
        for col in spec.get("non_negative", []):
            value = get(col, row)
            if value not in (None, "") and not (
                isinstance(value, (int, float)) and not isinstance(value, bool) and value >= 0
            ):
                problems.append(f"{key}: {col} must be a number >= 0 (got {value!r})")
        for col, ref in spec.get("values_in", {}).items():
            value = get(col, row)
            if value not in (None, "") and str(value).strip() not in lookup(ref):
                problems.append(f"{key}: {col} {value!r} not in {ref}")
        for col, ref in spec.get("ids_in", {}).items():
            ids = re.findall(fmt(profile, "req_id"), str(get(col, row) or ""))
            if not ids:
                problems.append(f"{key}: {col} has no ID")
            unknown = [i for i in ids if i not in lookup(ref)]
            if unknown:
                problems.append(f"{key}: {col} {unknown} not in {ref}")
        for col in spec.get("warn_empty", []):
            if get(col, row) in (None, ""):
                warns.append(f"{key}: {col}")
    if spec.get("unique_key"):
        dupes = sorted({k for k in keys if keys.count(k) > 1})
        if dupes:
            problems.append(f"duplicate keys {dupes}")
    r.check(bool(keys), f"{name}: has rows", "no rows")
    r.check(not problems, f"{name}: rows complete and consistent", "; ".join(problems[:8]))
    if warns:
        r.warn(f"{name}: optional cells empty", ", ".join(warns[:8]))


def evaluate(path, profile, r):
    """Read computed values for the profile's cells from a recalculated copy.

    LibreOffice is the converter's only formula engine. The copy is private:
    the published workbook stays exactly what the caller produced.
    """
    cells = [e["cell"] for e in profile.get("evaluate", [])]
    cells += [
        spec["cell"]
        for spec in profile.get("exports", {}).values()
        if isinstance(spec, dict) and "cell" in spec
    ]
    if not cells:
        return {}
    try:
        recalculated = pack.recalculate_xlsx(Path(path))
    except pack.PackError as exc:
        r.fail("formulas evaluated", f"recalculation unavailable: {exc}")
        return {}
    try:
        from openpyxl import load_workbook

        wb = load_workbook(recalculated, data_only=True)
        out = {}
        for ref in cells:
            sheet, cell = split_ref(ref)
            value = wb[sheet][cell].value if sheet in wb.sheetnames else None
            try:
                out[ref] = float(value) if value is not None and not isinstance(value, bool) else None
            except (TypeError, ValueError):
                out[ref] = None
        return out
    finally:
        Path(recalculated).unlink(missing_ok=True)


# ---------------------------------------------------------------- printed block


def check_page_fit(profile, output_path, r, workdir):
    """Assert a named block stays on one printed page (one page wide and tall).

    The anchors name the first and last text of the block. A block that fits
    one page width but runs onto the next page fails exactly as the manual
    check did, and now the finding is machine-reported.
    """
    spec = profile.get("pagination", {}).get("block_fits_one_page")
    if not spec:
        return
    start_text, end_text = spec.get("start"), spec.get("end")
    if not isinstance(start_text, str) or not start_text or not isinstance(end_text, str) or not end_text:
        r.fail("pagination: block fits one page", "the profile needs non-empty start and end anchors")
        return
    max_pages = spec.get("max_pages", 1)
    if not isinstance(max_pages, int) or isinstance(max_pages, bool) or max_pages < 1:
        r.fail("pagination: block fits one page", "max_pages must be a positive integer")
        return
    try:
        pdf = pack.convert_pdf(Path(output_path), Path(workdir), "page-fit")
    except pack.PackError as exc:
        r.fail("pagination: block fits one page", f"PDF unavailable: {exc}")
        return
    try:
        import fitz

        with fitz.open(pdf) as doc:
            start = next((i for i, page in enumerate(doc) if page.search_for(start_text)), None)
            end = next(
                (i for i in range(len(doc) - 1, -1, -1) if doc[i].search_for(end_text)),
                None,
            )
            if start is None or end is None:
                missing = [text for text, page in ((start_text, start), (end_text, end)) if page is None]
                r.fail("pagination: block fits one page", f"anchor text not found: {missing}")
                return
            if end < start:
                r.fail(
                    "pagination: block fits one page",
                    f"'{end_text}' appears before '{start_text}' (end page {end + 1} < start page {start + 1})",
                )
                return
            span = end - start + 1
            r.check(
                span <= max_pages,
                "pagination: block fits one page",
                f"'{start_text}'..'{end_text}' spans {span} page(s), max {max_pages}",
            )
    finally:
        pdf.unlink(missing_ok=True)


# ---------------------------------------------------------------- entry


def run(profile, template_path, output_path, source=None, workdir=None):
    r = Report()
    checker = {"docx": check_docx, "xlsx": check_xlsx}[profile["format"]]
    exports = checker(profile, template_path, output_path, r, source or {})
    if workdir is not None:
        check_page_fit(profile, output_path, r, workdir)
    return r, exports or {}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--profile", required=True)
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--manifest")
    parser.add_argument("--source-profile")
    parser.add_argument("--source-template")
    parser.add_argument("--source-output")
    parser.add_argument("--expected-total")
    parser.add_argument("--ids")
    args = parser.parse_args()

    try:
        profile = pack.load_profile(Path(args.profile))
        manifest = pack.load_manifest(Path(args.manifest) if args.manifest else None, profile)
        label = pack.summary_label(manifest)
        source = {}
        source_report = None
        if args.source_profile:
            if not (args.source_template and args.source_output):
                raise pack.PackError("--source-profile needs --source-template and --source-output")
            source_profile = pack.load_profile(Path(args.source_profile))
            source_report, exports = run(source_profile, args.source_template, args.source_output)
            source = exports
        if args.expected_total:
            source["local_total"] = args.expected_total
        if args.ids:
            source["in_scope_ids"] = [i.strip() for i in args.ids.split(",") if i.strip()]

        output_dir = Path(args.output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        r, _ = run(profile, args.template, args.output, source=source, workdir=output_dir)
        if source_report is not None:
            if source_report.fails:
                r.fail(
                    "source: reference document",
                    f"{source_report.fails} finding(s) in the source; see the source report",
                )
            declared = sorted((source_profile.get("exports") or {}).keys())
            missing = [key for key in declared if key not in source]
            if missing:
                r.fail("source: exports available", f"missing {missing}")

        lines = list(r.lines)
        if source_report is not None:
            lines = ["== source"] + source_report.lines + [source_report.summary("source")] + ["== output"] + lines
        summary = r.summary(label)
        text = "\n".join(lines + [summary]) + "\n"
        (output_dir / "report.txt").write_text(text, encoding="utf-8")
        (output_dir / "report.json").write_text(
            json.dumps(
                {
                    "schema": SCHEMA,
                    "passed": r.fails == 0,
                    "summary": summary,
                    "label": label,
                    "counts": {"pass": r.passes(), "warn": r.warns, "fail": r.fails},
                    "manifest": manifest,
                    "lines": lines,
                },
                indent=2,
                sort_keys=True,
            )
            + "\n",
            encoding="utf-8",
        )
        print(summary)
        print(f"report: {r.passes()} pass, {r.warns} warn, {r.fails} fail")
        return 0
    except pack.PackError as exc:
        print(f"Check refused: {exc}", file=sys.stderr)
        return 2
    except Exception as exc:  # an unreadable document is a refusal, not a pass
        print(f"Check refused: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
